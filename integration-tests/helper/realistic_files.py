"""Files that look like what people actually upload, built at test time.

The rest of the suite indexes files of 12 bytes to about 1 KB with plain ASCII
names. Real knowledge bases hold multi-megabyte exports, spreadsheets with
several sheets, long Word documents and PDFs, non-English names and folders
nested several levels deep. Each of those takes a different path through
parsing and chunking, and none of them was exercised.

Everything here is generated from a fixed seed, so a run is repeatable and
nothing large is committed to the repository. Each file carries one or more
made-up words that appear nowhere else (``tokens``). Searching for one has
exactly one right answer, which is what lets a test say *which* file came
back rather than merely that something did. A file with several tokens puts
each in a different place -- one per sheet, the last page -- so a parser that
reads only the first sheet or the first page is caught.
"""

from __future__ import annotations

import csv
import datetime
import io
import os
import random
import re
from dataclasses import dataclass

DEFAULT_LARGE_FILE_MB = 3.0
LARGE_FILE_MB_ENV = "PIPESHUB_IT_LARGE_FILE_MB"
# Office formats stamp creation times into the file; pinning them keeps the
# bytes identical between runs.
_FIXED_TIME = datetime.datetime(2026, 1, 1, 0, 0, 0)

# Ordinary words only. Tokens are built so they can never be spelled from
# these, which keeps every token's only source the one place it was put.
_VOCABULARY = (
    "account agenda allocation analysis approval archive audit backlog budget "
    "calendar capacity change client committee contract cost customer dashboard "
    "deadline decision delivery department deployment design draft estimate "
    "feedback finance forecast goal handover headcount incident invoice issue "
    "launch ledger meeting milestone migration minutes network onboarding "
    "outage owner partner payroll pipeline plan policy priority procurement "
    "project proposal quarter release renewal report request review risk "
    "roadmap schedule security service signoff sprint stakeholder summary "
    "supplier support target team template timeline training travel update "
    "vendor version workflow"
).split()


@dataclass(frozen=True)
class RealisticFile:
    """One file to upload.

    ``upload_path`` is what the knowledge-base page sends as the file's path.
    A path with folders in it is how a folder upload creates those folders.
    """

    slug: str
    name: str
    body: bytes
    mimetype: str
    tokens: tuple[str, ...]
    upload_path: str

    @property
    def size(self) -> int:
        return len(self.body)


def _sentence(rng: random.Random) -> str:
    words = [rng.choice(_VOCABULARY) for _ in range(rng.randint(8, 16))]
    return " ".join(words).capitalize() + "."


def _paragraph(rng: random.Random, sentences: int = 5) -> str:
    return " ".join(_sentence(rng) for _ in range(sentences))


def unicode_markdown() -> RealisticFile:
    """Several scripts in the name and the body, the way a global team writes."""
    token = "vexmorakli3381"
    name = "Q3 résumé — 季度报告 📊.md"
    body = "\n\n".join([
        "# Résumé du trimestre — 季度报告 — ملخص الربع — त्रैमासिक सारांश 📊",
        "Les équipes de Zürich et de São Paulo ont livré la migration à temps. "
        "Coördinatie met het team in Ålesund verliep soepel.",
        "上海团队完成了客户数据迁移，预算控制在计划之内。",
        "أنهى فريق دبي مراجعة الأمان قبل الموعد النهائي.",
        "मुंबई टीम ने ग्राहक सहायता प्रक्रिया को अपडेट किया।",
        f"Reference code for this review: {token} ✅🚀",
        "Emoji status legend: ✅ done, ⏳ waiting, ❌ blocked, 🔥 urgent.",
    ]).encode("utf-8")
    return RealisticFile(
        slug="unicode",
        name=name,
        body=body,
        mimetype="text/markdown",
        tokens=(token,),
        upload_path=name,
    )


def large_text(size_mb: float | None = None) -> RealisticFile:
    """A multi-megabyte export with one distinctive line well past the start.

    The token sits about 60% of the way in. A parser or chunker that truncates
    a long file, or a size limit hit part-way through indexing, leaves it
    unfindable.
    """
    if size_mb is None:
        size_mb = float(os.getenv(LARGE_FILE_MB_ENV, DEFAULT_LARGE_FILE_MB))
    target = int(size_mb * 1024 * 1024)
    token = "quiltharbex5127"
    rng = random.Random(5127)

    parts: list[str] = []
    written = 0
    placed = False
    section = 0
    while written < target:
        section += 1
        chunk = f"Section {section}\n\n{_paragraph(rng, 6)}\n\n"
        if not placed and written >= target * 0.6:
            chunk += f"Escalation reference {token} was raised in this section.\n\n"
            placed = True
        parts.append(chunk)
        written += len(chunk)
    return RealisticFile(
        slug="large",
        name="operations-log-export.txt",
        body="".join(parts).encode("utf-8"),
        mimetype="text/plain",
        tokens=(token,),
        upload_path="operations-log-export.txt",
    )


def multi_sheet_xlsx() -> RealisticFile:
    """A workbook with three sheets, a distinctive value deep in each one."""
    from openpyxl import Workbook  # noqa: PLC0415 - only needed when generating

    tokens = ("brindlecove2204", "sprocketfen7719", "marrowgale4406")
    rng = random.Random(2204)
    workbook = Workbook()
    workbook.properties.created = _FIXED_TIME
    workbook.properties.modified = _FIXED_TIME
    sheet_names = ("Revenue", "Headcount", "Vendors")
    for index, (sheet_name, token) in enumerate(zip(sheet_names, tokens, strict=True)):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = sheet_name
        sheet.append(["Region", "Owner", "Quarter", "Amount", "Notes"])
        for row in range(1, 61):
            note = f"Audit tag {token}" if row == 45 else _sentence(rng)
            sheet.append([
                rng.choice(["EMEA", "APAC", "AMER", "LATAM"]),
                rng.choice(["Amélie", "Søren", "李娜", "Priya", "Oluwaseun"]),
                f"2026-Q{rng.randint(1, 4)}",
                round(rng.uniform(100, 100000), 2),
                note,
            ])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return RealisticFile(
        slug="xlsx",
        name="finance-workbook-2026.xlsx",
        body=_normalise_zip(buffer.getvalue()),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        tokens=tokens,
        upload_path="finance-workbook-2026.xlsx",
    )


def multi_page_docx() -> RealisticFile:
    """A long Word document: headings, a table, page breaks, token on the last page."""
    from docx import Document  # noqa: PLC0415 - only needed when generating
    from docx.enum.text import WD_BREAK  # noqa: PLC0415

    token = "fennelquarz8830"
    rng = random.Random(8830)
    document = Document()
    document.core_properties.created = _FIXED_TIME
    document.core_properties.modified = _FIXED_TIME
    pages = 6
    for page in range(1, pages + 1):
        document.add_heading(f"Chapter {page}: {rng.choice(_VOCABULARY).title()} review", level=1)
        for _ in range(4):
            document.add_paragraph(_paragraph(rng, 5))
        if page == 2:
            table = document.add_table(rows=4, cols=3)
            for r, row in enumerate(table.rows):
                for c, cell in enumerate(row.cells):
                    cell.text = "Header" if r == 0 else f"{rng.choice(_VOCABULARY)} {r}.{c}"
        if page == pages:
            document.add_paragraph(f"Closing action item {token} is owned by the platform team.")
        else:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    buffer = io.BytesIO()
    document.save(buffer)
    return RealisticFile(
        slug="docx",
        name="annual-operations-review.docx",
        body=_normalise_zip(buffer.getvalue()),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        tokens=(token,),
        upload_path="annual-operations-review.docx",
    )


def multi_page_pdf() -> RealisticFile:
    """A text PDF of several pages, token on the last page only."""
    token = "harrowmint6652"
    rng = random.Random(6652)
    pages: list[list[str]] = []
    for page in range(1, 6):
        lines = [f"Board pack page {page}"]
        lines += [_sentence(rng) for _ in range(20)]
        pages.append(lines)
    pages[-1].append(f"Follow-up reference {token} closes this pack.")
    return RealisticFile(
        slug="pdf",
        name="board-pack.pdf",
        body=build_text_pdf(pages),
        mimetype="application/pdf",
        tokens=(token,),
        upload_path="board-pack.pdf",
    )


def messy_csv() -> RealisticFile:
    """A CSV with quoted commas, a line break inside a cell, and non-ASCII names."""
    token = "gravelpont9014"
    rng = random.Random(9014)
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(["id", "customer", "address", "notes"])
    for row in range(1, 200):
        writer.writerow([
            row,
            rng.choice(["Müller, Jürgen", "O'Brien & Sons", "Zoë Ångström", "株式会社テスト"]),
            f"{rng.randint(1, 999)} Main St, Suite {rng.randint(1, 50)}",
            "Line one\nline two, with a comma" if row % 17 == 0 else _sentence(rng),
        ])
    writer.writerow([200, "Final Customer", "1 Last Rd", f"Renewal code {token}"])
    return RealisticFile(
        slug="csv",
        name="customers-export.csv",
        body=buffer.getvalue().encode("utf-8"),
        mimetype="text/csv",
        tokens=(token,),
        upload_path="customers-export.csv",
    )


def deeply_nested_markdown() -> RealisticFile:
    """A file five folders down, uploaded the way a folder upload sends it."""
    token = "lanternveck2297"
    name = "incident-review.md"
    body = (
        "# Incident review\n\n"
        "The failover completed in under four minutes and no data was lost.\n\n"
        f"Tracking reference {token}.\n"
    ).encode("utf-8")
    return RealisticFile(
        slug="nested",
        name=name,
        body=body,
        mimetype="text/markdown",
        tokens=(token,),
        upload_path=f"Ingeniería/Plataforma/Runbooks/2026/Q3 回顾/{name}",
    )


# The corpus's slugs, known without generating anything, so tests can be
# parametrised by them at collection time.
REALISTIC_SLUGS = ("unicode", "large", "xlsx", "docx", "pdf", "csv", "nested")


def realistic_corpus(large_file_mb: float | None = None) -> list[RealisticFile]:
    """Every realistic file, in upload order (``REALISTIC_SLUGS``)."""
    return [
        unicode_markdown(),
        large_text(large_file_mb),
        multi_sheet_xlsx(),
        multi_page_docx(),
        multi_page_pdf(),
        messy_csv(),
        deeply_nested_markdown(),
    ]


def _normalise_zip(data: bytes) -> bytes:
    """Rewrite an Office zip with fixed timestamps so the bytes are repeatable."""
    import zipfile  # noqa: PLC0415

    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl stamps the save time here whatever the workbook says.
                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2026-01-01T00:00:00Z\g<2>",
                    content,
                )
            target.writestr(fixed, content)
    return out.getvalue()


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_text_pdf(pages: list[list[str]]) -> bytes:
    """A minimal valid PDF with one text line per entry, one page per list.

    No PDF library is a dependency of the suite, and a real text layer is the
    point: an image-only PDF would need OCR and test something else.
    Latin-1 text only, which is what the built-in Helvetica font can show.
    """
    objects: list[bytes] = []
    page_count = len(pages)
    font_id = 3
    first_page_id = 4
    kids = " ".join(f"{first_page_id + 2 * i} 0 R" for i in range(page_count))

    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for index, lines in enumerate(pages):
        content_id = first_page_id + 2 * index + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
            ).encode()
        )
        ops = ["BT", "/F1 10 Tf", "14 TL", "50 750 Td"]
        for line in lines:
            ops.append(f"({_pdf_escape(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    )
    return out.getvalue()
