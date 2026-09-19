"""The realistic test files say what they claim to (no live services).

The permission-matrix and realistic-data suites decide which file a search hit
came from by a token put inside it. If a token were missing, unreadable, or
present in two files, those suites would report the wrong thing, so the
generators are checked here by reading each file back with an ordinary library.
"""

from __future__ import annotations

import csv
import io

import pytest

from helper.realistic_files import (
    REALISTIC_SLUGS,
    build_text_pdf,
    large_text,
    multi_page_docx,
    multi_page_pdf,
    multi_sheet_xlsx,
    realistic_corpus,
)

pytestmark = pytest.mark.unit


def _text_of(file) -> str:
    """What a parser would read out of the file."""
    if file.slug == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(file.body), read_only=True)
        return "\n".join(
            str(cell) for sheet in workbook for row in sheet.iter_rows(values_only=True) for cell in row
        )
    if file.slug == "docx":
        from docx import Document

        return "\n".join(p.text for p in Document(io.BytesIO(file.body)).paragraphs)
    if file.slug == "pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(file.body)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    return file.body.decode("utf-8")


@pytest.fixture(scope="module")
def corpus():
    return realistic_corpus(large_file_mb=0.5)


def test_every_token_is_readable_from_its_own_file(corpus) -> None:
    for file in corpus:
        text = _text_of(file)
        for token in file.tokens:
            assert token in text, f"{file.slug}: token {token} not readable from the file"


def test_no_token_appears_in_any_other_file(corpus) -> None:
    texts = {file.slug: _text_of(file) for file in corpus}
    for file in corpus:
        for token in file.tokens:
            holders = [slug for slug, text in texts.items() if token in text]
            assert holders == [file.slug], f"token {token} found in {holders}"


def test_declared_slugs_match_the_corpus(corpus) -> None:
    assert tuple(file.slug for file in corpus) == REALISTIC_SLUGS


def test_slugs_and_tokens_are_unique(corpus) -> None:
    slugs = [file.slug for file in corpus]
    tokens = [token for file in corpus for token in file.tokens]
    assert len(set(slugs)) == len(slugs)
    assert len(set(tokens)) == len(tokens)


def test_generation_is_repeatable(corpus) -> None:
    again = realistic_corpus(large_file_mb=0.5)
    for first, second in zip(corpus, again, strict=True):
        assert first.body == second.body, f"{first.slug} differs between two generations"


def test_each_sheet_carries_its_own_token() -> None:
    from openpyxl import load_workbook

    file = multi_sheet_xlsx()
    workbook = load_workbook(io.BytesIO(file.body), read_only=True)
    assert len(workbook.sheetnames) == len(file.tokens) == 3
    for sheet_name, token in zip(workbook.sheetnames, file.tokens, strict=True):
        cells = [str(c) for row in workbook[sheet_name].iter_rows(values_only=True) for c in row]
        assert any(token in cell for cell in cells), f"{token} missing from sheet {sheet_name}"


def test_pdf_token_is_on_the_last_page_only() -> None:
    import pdfplumber

    file = multi_page_pdf()
    with pdfplumber.open(io.BytesIO(file.body)) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    assert len(pages) >= 5
    assert file.tokens[0] in pages[-1]
    assert not any(file.tokens[0] in page for page in pages[:-1])


def test_docx_token_comes_after_several_page_breaks() -> None:
    from docx import Document

    file = multi_page_docx()
    xml = Document(io.BytesIO(file.body)).element.xml
    before_token = xml.split(file.tokens[0])[0]
    assert before_token.count('w:type="page"') >= 5


def test_large_file_reaches_its_size_and_hides_the_token_past_halfway() -> None:
    file = large_text(size_mb=1.0)
    assert file.size >= 1024 * 1024
    position = file.body.find(file.tokens[0].encode())
    assert position > file.size * 0.5
    assert file.body.count(file.tokens[0].encode()) == 1


def test_csv_parses_with_quoted_commas_and_line_breaks(corpus) -> None:
    file = next(f for f in corpus if f.slug == "csv")
    rows = list(csv.reader(io.StringIO(file.body.decode("utf-8"))))
    assert all(len(row) == 4 for row in rows)
    assert any("\n" in row[3] for row in rows)
    assert any("," in row[1] for row in rows[1:])


def test_names_and_folders_are_not_plain_ascii(corpus) -> None:
    by_slug = {file.slug: file for file in corpus}
    assert not by_slug["unicode"].name.isascii()
    nested = by_slug["nested"].upload_path.split("/")
    assert len(nested) >= 6, "the file should sit at least five folders down"
    assert not "/".join(nested[:-1]).isascii()


def test_pdf_escapes_parentheses_and_backslashes() -> None:
    import pdfplumber

    body = build_text_pdf([["a (bracketed) value", "a back\\slash"]])
    with pdfplumber.open(io.BytesIO(body)) as pdf:
        text = pdf.pages[0].extract_text()
    assert "a (bracketed) value" in text
    assert "a back\\slash" in text
