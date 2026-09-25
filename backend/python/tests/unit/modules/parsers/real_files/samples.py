"""Builders for small real sample files, and helpers to read back what a parser produced."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from app.models.blocks import BlocksContainer, BlockType

if TYPE_CHECKING:
    from collections.abc import Iterable, Sequence


def make_docx(
    heading: str,
    paragraphs: Sequence[str],
    table: Sequence[Sequence[str]] = (),
    bullets: Sequence[str] = (),
    closing: str | None = None,
) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(heading, 1)
    for text in paragraphs:
        document.add_paragraph(text)
    for text in bullets:
        document.add_paragraph(text, style="List Bullet")
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    if closing:
        document.add_paragraph(closing)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx(slides: Iterable[tuple[str, str]], table: Sequence[Sequence[str]] = ()) -> bytes:
    import pptx
    from pptx.util import Inches

    presentation = pptx.Presentation()
    for title, body in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    if table:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Table slide"
        shape = slide.shapes.add_table(
            len(table), len(table[0]), Inches(1), Inches(2), Inches(6), Inches(3)
        )
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                shape.table.cell(r, c).text = value
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets: dict[str, Sequence[Sequence[object]]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pdf(pages: Sequence[Sequence[str]], encrypt: str | None = None) -> bytes:
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, encrypt=encrypt)
    for lines in pages:
        y = 750
        for line in lines:
            pdf.drawString(72, y, line)
            y -= 18
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_png(width: int = 8, height: int = 6) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


def all_text(container: BlocksContainer) -> str:
    """Every piece of text a container would send to the index, in block order."""
    parts: list[str] = []
    for block in container.blocks:
        if block.type == BlockType.TABLE_ROW and isinstance(block.data, dict):
            parts.append(str(block.data.get("row_natural_language_text", "")))
        elif isinstance(block.data, str):
            parts.append(block.data)
    return "\n".join(parts)


def table_rows(container: BlocksContainer) -> list[str]:
    return [
        str(block.data.get("row_natural_language_text", ""))
        for block in container.blocks
        if block.type == BlockType.TABLE_ROW and isinstance(block.data, dict)
    ]
